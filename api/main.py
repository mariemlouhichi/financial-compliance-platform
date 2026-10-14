r"""
api/main.py — Backend FastAPI du système AML/KYC (Point 6)
==============================================================================
Expose le pipeline complet (les 8 étapes réelles du projet) au frontend React :

  [1] Level 1        → décision déterministe + signaux
  [2] Client Scoring → profil structurel (PEP/KYC/FATCA)
  [3] Planning       → questions réglementaires
  [4] Retrieval      → chunks BCT/FATF (Weaviate + Neo4j)
  [5] Compliance     → décision neuro-symbolique + lois violées/respectées
  [6] ML             → proba + SHAP (signal indicatif, divergence)
  [7] SAR            → rapport complet (markdown + JSON)
  [8] Workflow       → état (APPROVED/BLOCKED/PENDING_*) + file superviseur

ENDPOINTS
  POST /api/analyze              lance l'analyse (asynchrone) → {tx_id}
  GET  /api/result/{tx_id}       statut + sorties de CHAQUE étape + SAR
  GET  /api/sar/{tx_id}/markdown le rapport SAR brut (texte)
  GET  /api/transactions         historique des analyses (résumés)
  POST /api/validate/{tx_id}     décision superviseur (APPROVE/REJECT/ANNOTATE)
  GET  /api/stats                KPIs pour le dashboard
  GET  /api/clients/{client_id}  profil KYC (pour pré-remplir le formulaire)
  GET  /api/demo-transactions    les 6 transactions de test (simulateur)
  GET  /api/regulations/search   recherche dans le corpus réglementaire
  GET  /api/health               statut des composants

LANCEMENT (l'init charge les modèles — ~30-60s au premier appel retrieval) :
  .\venv\Scripts\python.exe -m uvicorn api.main:app --port 8001
  Swagger : http://localhost:8001/docs   (le front Vite appelle :8001)

NOTE : port 8001 (8000 = ml_api.py Swagger ML).
"""

import os
import sys
import json
import random
import time
import asyncio
import threading
import traceback
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from typing import Optional

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))
sys.path.insert(0, str(BASE / "level1_rules"))

from fastapi import FastAPI, HTTPException, Header
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import PlainTextResponse, FileResponse, StreamingResponse
from pydantic import BaseModel, Field

# ── Composants du pipeline (tous existants) ─────────────────────────────
from level1_rules.level1_orchestrator import Level1Orchestrator
from client_scoring_agent import ClientScoringAgent, client_score_to_signals
from clients_db_loader import ClientDBLoader
from client_history_store import get_history_store
from planning_agent import PlanningAgent
from retrieval_agent import RetrievalPipeline
from compliance_agent import ComplianceAgent
from sar_generator import SARGenerator
from decision_workflow import DecisionWorkflow, JsonlNotificationSink

# Réutilise MLPredictor + helpers du test E2E (source unique de vérité)
from test_e2e_pipeline import (MLPredictor, l1_to_planning, infer_jurisdiction,
                                fatf_checklist, deduplicate_chunks)

# Authentification (comptes + confirmation email + tokens) — api/auth.py
from api.auth import (auth_router, init_db as auth_init_db, current_user,
                      get_user_raw_profile, get_user_email, send_decision_email,
                      send_weekly_report_email)
# Persistance des dossiers : PostgreSQL si disponible, sinon repli fichiers JSON
from api.store import get_store
# Planificateur des tâches périodiques (réentraînement, rapport hebdomadaire)
from api.scheduler import SCHEDULER, Job, FRIDAY, SUNDAY

try:
    from neo4j import GraphDatabase
    NEO4J_AVAILABLE = True
except ImportError:
    NEO4J_AVAILABLE = False


# ═══════════════════════════════════════════════════════════════════════
#  ÉTAT GLOBAL — composants (init unique) + registre des jobs
# ═══════════════════════════════════════════════════════════════════════

C: dict = {}                 # composants initialisés
JOBS: dict = {}              # tx_id → job {status, steps, sar, ...}
_init_lock = threading.Lock()
_counter_lock = threading.Lock()
_tx_counter = 0

# ── PERSISTANCE DES DOSSIERS ─────────────────────────────────────────────
#  JOBS est en mémoire → sans persistance, chaque redémarrage d'uvicorn
#  effaçait TOUT l'historique (l'utilisateur ne retrouvait plus ses dossiers
#  en se reconnectant). Chaque dossier est désormais persisté via une couche
#  d'accès (api/store.py) : PostgreSQL si disponible, sinon repli fichiers
#  JSON dans output/jobs/{tx_id}.json. Le backend est choisi une fois au
#  chargement (STORE) et ne bloque jamais l'API.
JOBS_DIR = Path(__file__).resolve().parent.parent / "output" / "jobs"
STORE = get_store(JOBS_DIR)


# ── BUS D'ÉVÉNEMENTS TEMPS RÉEL (SSE — Tâche 2) ──────────────────────────
#  Remplace le polling : le serveur POUSSE une notification aux clients dès
#  qu'un dossier change (analyse terminée, décision superviseur). Le pipeline
#  tourne dans un THREAD → publish() doit être thread-safe : on planifie la
#  mise en file sur la boucle asyncio via call_soon_threadsafe.
class EventBus:
    def __init__(self):
        self._subs: set = set()
        self._loop: Optional[asyncio.AbstractEventLoop] = None

    async def subscribe(self) -> asyncio.Queue:
        self._loop = asyncio.get_running_loop()
        q: asyncio.Queue = asyncio.Queue(maxsize=100)
        self._subs.add(q)
        return q

    def unsubscribe(self, q: asyncio.Queue):
        self._subs.discard(q)

    def publish(self, event: dict):
        """Diffuse un événement à tous les abonnés — appelable depuis n'importe
        quel thread (jamais bloquant, jamais levant)."""
        loop = self._loop
        if loop is None or not self._subs:
            return
        for q in list(self._subs):
            try:
                loop.call_soon_threadsafe(_bus_put, q, event)
            except RuntimeError:
                pass


def _bus_put(q: asyncio.Queue, event: dict):
    try:
        q.put_nowait(event)
    except asyncio.QueueFull:
        pass


BUS = EventBus()


def _publish_job_event(job: dict, kind: str):
    """Publie l'événement de changement d'un dossier (charge utile minimale)."""
    try:
        summary = job.get("summary") or {}
        BUS.publish({
            "type": kind,                      # 'job' | 'decision'
            "tx_id": job.get("tx_id"),
            "owner": job.get("owner"),
            "status": job.get("status"),
            "decision": summary.get("ca_decision"),
            "state": summary.get("state"),
            "at": datetime.now().isoformat(),
        })
    except Exception:
        pass


def _persist_job(job: dict):
    """Persiste le dossier (backend PostgreSQL ou fichiers) — jamais bloquant."""
    STORE.save(job)


def _load_jobs():
    """Recharge l'historique au démarrage + réaligne le compteur d'IDs
    (sinon les nouveaux TX-YYYYMMDD-0001 écraseraient ceux du même jour)."""
    global _tx_counter
    n = 0
    for tx_id, job in STORE.load_all().items():
        if job.get("status") == "RUNNING":     # interrompu par un restart
            job["status"] = "ERROR"
            job["error"] = "Analyse interrompue par un redémarrage du serveur"
        JOBS[tx_id] = job
        n += 1
    today = datetime.now().strftime("%Y%m%d")
    mx = 0
    for tx_id in JOBS:
        parts = tx_id.split("-")
        if len(parts) == 3 and parts[1] == today:
            try:
                mx = max(mx, int(parts[2]))
            except ValueError:
                pass
    with _counter_lock:
        _tx_counter = max(_tx_counter, mx)
    if n:
        print(f"  [API] 🗂️ {n} dossier(s) rechargé(s) depuis {STORE.kind}")


def _next_tx_id() -> str:
    global _tx_counter
    with _counter_lock:
        _tx_counter += 1
        return f"TX-{datetime.now().strftime('%Y%m%d')}-{_tx_counter:04d}"


def init_components():
    """Initialise tous les composants une seule fois (thread-safe)."""
    with _init_lock:
        if C.get("ready"):
            return
        print("  [API] Initialisation des composants...")
        C["l1"] = Level1Orchestrator()
        from level1_rules.risk_country_checker import RiskCountryChecker
        C["client"] = ClientScoringAgent(country_checker=RiskCountryChecker())
        C["history"] = get_history_store(verbose=True)
        C["clients_db"] = ClientDBLoader(str(BASE / "clients_db.json"), verbose=True)
        C["retriever"] = RetrievalPipeline(verbose=False)
        C["planner"] = PlanningAgent(verbose=False, retriever=C["retriever"])
        neo4j_driver = None
        if NEO4J_AVAILABLE:
            try:
                # Paramétrable par variables d'environnement : en conteneur,
                # Neo4j se joint par son nom de service, pas par localhost.
                neo4j_driver = GraphDatabase.driver(
                    os.getenv("NEO4J_URI", "bolt://localhost:7687"),
                    auth=(os.getenv("NEO4J_USER", "neo4j"),
                          os.getenv("NEO4J_PASSWORD", "amlkyc2025")))
                neo4j_driver.verify_connectivity()
            except Exception:
                neo4j_driver = None
        C["neo4j"] = neo4j_driver
        C["compliance"] = ComplianceAgent(neo4j_driver=neo4j_driver, verbose=False)
        C["ml"] = MLPredictor(prefix="aml_ml_demo", history_store=C["history"])
        C["sar"] = SARGenerator()
        C["workflow"] = DecisionWorkflow(notifier=JsonlNotificationSink(
            path=str(BASE / "output" / "notifications.jsonl")), verbose=False)
        C["ready"] = True
        print("  [API] ✅ Composants prêts")


# ═══════════════════════════════════════════════════════════════════════
#  SÉRIALISATION (dataclasses → JSON pour le front)
# ═══════════════════════════════════════════════════════════════════════

def _rule_to_dict(r) -> dict:
    _thr = getattr(r, "threshold", None)
    return {
        "rule_text": (getattr(r, "rule_text", "") or "")[:250],
        "source_doc": getattr(r, "source_doc", ""),
        "rule_ref": getattr(r, "rule_ref", ""),
        "fatf_rec": getattr(r, "fatf_rec", ""),
        "severity": getattr(r, "severity", ""),
        "category": getattr(r, "category", ""),
        "provenance": getattr(r, "provenance", ""),
        "rule_status": getattr(r, "rule_status", ""),
        "violated": bool(getattr(r, "violated", False)),
        "mitigated": bool(getattr(r, "mitigated", False)),
        "confidence": round(float(getattr(r, "confidence", 0) or 0), 3),
        "cross_validated": bool(getattr(r, "cross_validated", False)),
        "cross_source_validated": bool(getattr(r, "cross_source_validated", False)),
        # [Front] citation VERBATIM du texte de loi — preuve anti-hallucination.
        # Uniquement si ancrée mot-à-mot (snippet_verbatim), sinon chaîne vide.
        "source_snippet": ((getattr(r, "source_snippet", "") or "")[:220]
                           if getattr(r, "snippet_verbatim", False) else ""),
        "threshold": (float(_thr) if isinstance(_thr, (int, float)) else None),
        "threshold_unit": getattr(r, "threshold_unit", ""),
    }


def _chunk_to_dict(c) -> dict:
    return {
        "source_doc": getattr(c, "source_doc", ""),
        "article_ref": getattr(c, "article_ref", ""),
        "jurisdiction": getattr(c, "jurisdiction", ""),
        "chunk_role": getattr(c, "chunk_role", ""),
        "score": round(float(getattr(c, "score", 0) or 0), 4),
        "retrieval_path": getattr(c, "retrieval_path", ""),
        "text": (getattr(c, "text", "") or "")[:300],
    }


# ── Explication FRANÇAISE de chaque question du Planning (pour l'admin) ──
#  Le Planning trace la provenance (source) et le déclencheur (factor) de
#  chaque question ; on les traduit en une phrase lisible qui montre au jury
#  POURQUOI le système pose cette question.

_SIGNAL_FR = {
    "SANCTION_UN_INDIVIDUAL_SENDER":      "l'émetteur figure sur la liste de sanctions ONU",
    "SANCTION_UN_INDIVIDUAL_BENEFICIARY": "le bénéficiaire figure sur la liste de sanctions ONU",
    "SANCTION_UN_ENTITY_SENDER":          "l'entité émettrice est sanctionnée (ONU)",
    "SANCTION_PEP_SENDER":                "l'émetteur est une Personne Politiquement Exposée (FATF R12)",
    "SANCTION_PEP_BENEFICIARY":           "le bénéficiaire est une Personne Politiquement Exposée (FATF R12)",
    "COUNTRY_BLACK_DESTINATION_COUNTRY":  "le pays de destination est sur la LISTE NOIRE FATF",
    "COUNTRY_BLACK_SENDER_COUNTRY":       "le pays d'origine est sur la LISTE NOIRE FATF",
    "COUNTRY_GREY_DESTINATION_COUNTRY":   "le pays de destination est sur la liste grise FATF",
    "COUNTRY_GREY_SENDER_COUNTRY":        "le pays d'origine est sur la liste grise FATF",
    "THRESHOLD_STR":                      "le montant dépasse le seuil de déclaration de soupçon (BCT)",
    "THRESHOLD_EDD":                      "le montant impose une diligence renforcée (EDD)",
    "THRESHOLD_KYC":                      "le montant impose une identification KYC",
    "THRESHOLD_RECORD":                   "le montant impose une conservation d'enregistrement",
    "SMURFING_PATTERN":                   "un pattern de fractionnement (dépôts répétés sous le seuil) est détecté",
    "CLIENT_HIGH_RISK":                   "le profil client est à risque structurel ÉLEVÉ",
    "CLIENT_MEDIUM_RISK":                 "le profil client est à risque structurel MOYEN",
    "CLIENT_PRIOR_STR":                   "le client a des déclarations de soupçon antérieures",
    "CLIENT_KYC_INCOMPLETE":              "le dossier KYC du client est incomplet",
}

_FACTOR_FR = {
    "no_justification":      "l'opération n'a PAS de justificatif économique",
    "new_bills":             "dépôt en billets neufs (indicateur de blanchiment)",
    "passport_cross_border": "identification par passeport sur une opération transfrontalière",
    "withdrawal=CASH":       "retrait du mandat en espèces",
    "new_receiver":          "premier virement vers ce bénéficiaire",
    "prev_rejected":         "une opération précédente du client a été rejetée",
    "tx_type":               "obligation de base pour ce type d'opération",
}


def _explain_question(source: str, factor: str, tx_type: str) -> str:
    """Une phrase française expliquant pourquoi le Planning pose cette question."""
    f = factor or ""
    if source == "template":
        return f"Vérification standard obligatoire pour toute opération {tx_type} (KYC, seuils, déclaration)."
    if source == "signal":
        why = _SIGNAL_FR.get(f, f or "signal de risque Level 1")
        return f"Déclenchée par le Niveau 1 : {why}."
    if source == "context":
        if f.startswith("motive="):
            return f"Le motif déclaré ({f.split('=', 1)[1]}) est vague ou insuffisant — justification économique à vérifier."
        if f.startswith("hour="):
            return f"Opération à heure atypique ({f.split('=', 1)[1]}) — surveillance hors heures ouvrées."
        if f.startswith("just_below="):
            return "Montant JUSTE SOUS le seuil réglementaire — recherche d'un contournement (structuring)."
        if f.startswith("round="):
            return "Montant rond — indicateur faible de blanchiment, vérification des obligations."
        if f.startswith("intermediary="):
            return f"Banque intermédiaire ({f.split('=', 1)[1]}) — screening correspondant bancaire (FATF R13)."
        if f.startswith("operator="):
            return f"Opérateur de transfert {f.split('=', 1)[1]} — obligations MVTS."
        return f"Analyse du contexte : {_FACTOR_FR.get(f, f)}."
    if source == "history":
        if f.startswith("prev_strs="):
            return f"Historique : {f.split('=', 1)[1]} déclaration(s) de soupçon antérieure(s) — vigilance continue."
        if f.startswith("velocity="):
            return f"Historique : {f.split('=', 1)[1]} opérations en 24h — fréquence inhabituelle."
        if f.startswith("new_country="):
            return f"Historique : première opération vers {f.split('=', 1)[1]} — changement de comportement."
        if f.startswith("new_account="):
            return f"Historique : compte récent ({f.split('=', 1)[1]}) avec montant important."
        return f"Analyse de l'historique client : {_FACTOR_FR.get(f, f)}."
    if source == "llm_reason":
        return f"Raisonnement LLM (pattern multi-facteurs) : {f}" if f else \
            "Ajoutée par le raisonnement LLM — pattern que les règles fixes ne voient pas."
    return "Question réglementaire générée par le Planning Agent."


# ── Synthèse des réponses réglementaires (LLM sous contrainte) ───────────
#  Pour chaque question du Planning, produit UNE réponse lisible en français
#  citant ses sources — les extraits bruts (chunks) restent disponibles
#  derrière « détail » dans l'UI. UN SEUL appel LLM groupé (temp=0 → cache).
#  Repli déterministe : liste des sources, jamais d'échec bloquant.

def _synthesize_answers(by_question: list, tx_type: str):
    # Repli déterministe posé d'abord (remplacé si le LLM répond)
    for b in by_question:
        srcs = [f"{a['source_doc'].replace('.pdf', '')} {a['article_ref']}".strip()
                for a in b.get("answers", []) if a.get("source_doc")]
        b["synthesis"] = ""
        b["sources"] = list(dict.fromkeys(srcs))[:4]

    try:
        from llm_client import get_llm
        llm = get_llm()
    except Exception:
        llm = None
    if llm is None or not by_question:
        return

    blocks = []
    for i, b in enumerate(by_question, 1):
        extraits = "\n".join(
            f"  [{a['source_doc'].replace('.pdf', '')} · {a['article_ref'] or '—'}] "
            f"{a['text'][:220]}"
            for a in b.get("answers", [])[:4]) or "  (aucun extrait)"
        blocks.append(f"Q{i}: {b['question']}\nEXTRAITS:\n{extraits}")

    prompt = (
        "Tu es analyste conformité AML/KYC dans une banque tunisienne. Pour "
        "CHAQUE question ci-dessous, rédige en FRANÇAIS une réponse de 1 à 2 "
        "phrases qui synthétise ce que disent les EXTRAITS réglementaires "
        f"fournis (transaction de type {tx_type}).\n"
        "CONTRAINTES STRICTES :\n"
        "- Utilise UNIQUEMENT les extraits fournis pour chaque question. "
        "N'invente aucune obligation, aucun seuil, aucun article.\n"
        "- Cite la ou les sources entre parenthèses, ex. (BCT 2017-08 Art. 14).\n"
        "- Si les extraits ne répondent pas à la question, écris exactement : "
        "\"Les extraits trouvés ne répondent pas directement à cette question.\"\n\n"
        + "\n\n".join(blocks)
        + "\n\n## Sortie (JSON strict, sans markdown) : "
        '{"answers": ["réponse Q1", "réponse Q2", ...]} '
        f"— exactement {len(by_question)} réponses, dans l'ordre."
    )
    try:
        import re as _re
        resp = llm.chat.completions.create(
            model=None, messages=[{"role": "user", "content": prompt}],
            temperature=0.0, max_tokens=180 * len(by_question))
        raw = (resp.choices[0].message.content or "").strip()
        raw = _re.sub(r"```json\s*|```", "", raw)
        answers = json.loads(raw).get("answers", [])
        for b, a in zip(by_question, answers):
            if isinstance(a, str) and 20 <= len(a) <= 800:
                b["synthesis"] = a.strip()
    except Exception as e:
        # Repli silencieux — les sources déterministes restent affichées
        print(f"  [API] synthèse LLM indisponible ({str(e)[:80]}) — repli sources")


# ═══════════════════════════════════════════════════════════════════════
#  LE PIPELINE — exécute les 8 étapes et capture chaque sortie
# ═══════════════════════════════════════════════════════════════════════

def execute_pipeline(tx: dict, tx_id: str):
    """Version API de run_pipeline() : mêmes étapes, sorties JSON par étape."""
    job = JOBS[tx_id]
    steps = job["steps"]

    def done(name, t0, data):
        data["duration_ms"] = int((time.time() - t0) * 1000)
        steps[name] = data
        job["current_step"] = name

    try:
        # ── [0] JUSTIFICATIFS joints à la soumission (intégré au pipeline) ──
        #  Chaque document est extrait (docling/pypdf) puis analysé (vérifs
        #  déterministes + LLM sous contrainte). Le verdict REMPLACE la case
        #  déclarative has_justification : recevable → True, tous rejetés →
        #  False → le Level 1 et le Compliance traitent la transaction comme
        #  NON justifiée, quel que soit ce que l'utilisateur a coché.
        files = job.get("_files") or []
        if files:
            results = []
            for f in files:
                try:
                    results.append(_analyze_one_justificatif(
                        f.get("filename", "document"), f.get("content_b64", ""),
                        tx, []))
                except ValueError as e:
                    results.append({"filename": f.get("filename", "document"),
                                    "conforme": False, "checks": {}, "llm": None,
                                    "error": str(e), "size_kb": 0,
                                    "uploaded_at": datetime.now().isoformat(),
                                    "extracted_chars": 0, "new_state": None})
            job["justificatifs"] = results
            job["justificatif"] = results[-1]
            tx["has_justification"] = any(r["conforme"] for r in results)
            job["transaction"]["has_justification"] = tx["has_justification"]
            job["_files"] = None   # libère la mémoire (base64)

        # ── Identification du client (nécessaire dès le Level 1) ─────
        sender_name = tx.get("sender_name", "")
        client_id = (sender_name.replace("Account-", "")
                     if "Account-" in sender_name
                     else (tx.get("client_id") or sender_name))

        # ── [1] LEVEL 1 ───────────────────────────────────────────────
        #  MÉMOIRE : on injecte l'historique Redis des dépôts récents (72h)
        #  → _detect_smurfing voit les transactions PRÉCÉDENTES du client.
        #  Sans cela, le fractionnement (3+ dépôts < 5 000 TND/72h) serait
        #  indétectable transaction par transaction.
        t0 = time.time()
        history_72h = []
        try:
            history_72h = C["history"].recent_deposits(client_id)
        except Exception as e:
            steps.setdefault("warnings", []).append(f"history: {e}")
        l1 = C["l1"].check_transaction({**tx, "customer_history": history_72h})
        done("level1", t0, {
            "decision": l1.decision,
            "violations": [str(v)[:150] for v in l1.violations],
            "risk_signals": list(l1.risk_signals),
            "history_deposits_72h": len(history_72h),
        })

        # ── [2] CLIENT SCORING ────────────────────────────────────────
        t0 = time.time()
        # PRIORITÉ aux coordonnées KYC de l'utilisateur connecté : si le
        # propriétaire du dossier a renseigné son profil (page « Mon profil »),
        # c'est LUI le client scoré — ses coordonnées alimentent le scoring.
        owner_raw = get_user_raw_profile(job.get("owner"))
        if owner_raw:
            client_id = job["owner"]
            known = True
            profile = C["clients_db"]._to_profile(owner_raw)
        else:
            known = C["clients_db"].exists(client_id)
            profile = C["clients_db"].get_profile(client_id, fallback_default=True)
        client_result, client_signals = None, []
        client_step = {"client_id": client_id, "known": bool(known)}
        if profile:
            client_result = C["client"].score_client(profile)
            client_signals = client_score_to_signals(client_result)
            client_step.update({
                "nationality": getattr(profile, "nationality", ""),
                "is_pep": bool(getattr(profile, "is_pep", False)),
                "kyc_complete": bool(getattr(profile, "kyc_complete", False)),
                "previous_strs": getattr(profile, "previous_strs", 0),
                "level": client_result.level,
                "score": float(client_result.score),
                "factors": dict(client_result.components or {}),
                "signals": client_signals,
                "fatca": getattr(getattr(client_result, "fatca", None),
                                 "obligation", "NONE"),
            })
        done("client_scoring", t0, client_step)

        # ── [3] PLANNING ──────────────────────────────────────────────
        t0 = time.time()
        planning_tx = l1_to_planning(tx)
        all_signals = list(dict.fromkeys(list(l1.risk_signals) + client_signals))
        plan = C["planner"].plan(planning_tx, all_signals, None)
        tx_type = tx.get("tx_type", "SWIFT")
        # Chaque question porte son POURQUOI en français (lisible par l'admin)
        factors = list(getattr(plan, "factors", [])) or [""] * len(plan.questions)
        questions_detail = [
            {"question": q, "priority": p, "source": s, "factor": f,
             "explanation": _explain_question(s, f, tx_type)}
            for q, p, s, f in zip(plan.questions, plan.priorities,
                                  plan.sources, factors)
        ]
        done("planning", t0, {
            "questions": list(plan.questions),
            "priorities": list(plan.priorities),
            "sources": list(plan.sources),
            "questions_detail": questions_detail,
            "llm_used": bool(getattr(plan, "llm_reasoning_used", False)),
        })

        # ── [4] RETRIEVAL — GRAPHE **ET** CORPUS combinés ─────────────
        #  Pour CHAQUE question, DEUX sources interrogées puis fusionnées :
        #    📗 le graphe Neo4j (règles GT + apprises, instantané)
        #    🔎 le corpus Weaviate (BM25 + vecteurs + RRF + reranker BGE)
        #  → la réponse cite les règles gouvernées ET les textes sources.
        t0 = time.time()
        jurisdiction = infer_jurisdiction(tx)
        all_chunks = []
        by_question = []
        for q in plan.questions:
            entry = {"question": q, "path": "GRAPH+CORPUS",
                     "graph_count": 0, "corpus_count": 0, "answers": []}
            try:
                g_rules = C["retriever"].graph_lookup(q, tx_type)
                entry["graph_count"] = len(g_rules)
            except Exception as e:
                g_rules = []
                steps.setdefault("warnings", []).append(f"graph_lookup: {e}")
            try:
                c_chunks = C["retriever"].retrieve(
                    query=q, tx_type=tx_type,
                    jurisdiction=jurisdiction, skip_cache=True, top_k=5)
                entry["corpus_count"] = len(c_chunks)
            except Exception as e:
                c_chunks = []
                entry["path"] = "GRAPH" if g_rules else "ERREUR"
                steps.setdefault("warnings", []).append(f"retrieval: {e}")
            merged = list(g_rules[:3]) + list(c_chunks)
            all_chunks.extend(merged)
            # marquage de provenance pour l'UI (📗 règle du graphe / 🔎 corpus)
            entry["answers"] = [
                {**_chunk_to_dict(c),
                 "from_graph": getattr(c, "retrieval_path", "") == "cache"}
                for c in merged[:4] if (getattr(c, "text", "") or "").strip()
            ]
            by_question.append(entry)
        deduped = deduplicate_chunks(all_chunks)

        # ── Synthèse LLM par question (UN appel groupé, sous contrainte) ──
        _synthesize_answers(by_question, tx_type)

        done("retrieval", t0, {
            "chunks_raw": len(all_chunks),
            "chunks_unique": len(deduped),
            "sources_count": len(set(c.source_doc for c in deduped)),
            "graph_answered": sum(1 for b in by_question if b["graph_count"] >= 1),
            "corpus_answered": sum(1 for b in by_question if b["corpus_count"] >= 1),
            "by_question": by_question,
            "top_chunks": [_chunk_to_dict(c) for c in deduped[:10]],
        })

        # ── [5] COMPLIANCE ────────────────────────────────────────────
        t0 = time.time()
        decision = C["compliance"].analyze(
            transaction=planning_tx,
            chunks=deduped,
            # all_signals = L1 ∪ client_signals (déjà fusionnés pour le Planning
            # à la step [3]) — les CLIENT_* sont mappés à poids 0 dans
            # SOFT_SIGNALS (compliance_agent.py) : visibles dans l'audit_trail,
            # zéro impact sur le score (déjà compté via client_profile→
            # client_level→CLIENT_RISK_CONTRIB/CLIENT_FLOOR, pas de double comptage).
            risk_signals=all_signals,
            fatf_checklist=fatf_checklist(tx.get("tx_type", "SWIFT")),
            level1_decision=l1.decision,
            # MODE PRODUCTION DUAL 3 ZONES : les faits durs (sanctions/listes
            # noires R6/R7/R19) restent symbolique-seul ; le chemin SÉMANTIQUE
            # (LLM + gate anti-hallucination) tourne sur les obligations
            # contextuelles → les règles validées (conf≥0.65) sont PERSISTÉES
            # dans Neo4j → le graphe apprend à chaque transaction (Step 0).
            dual_mode=True,
            client_profile=client_result)
        rc = decision.risk_components
        dstats = dict(getattr(decision, "stats", {}) or {})
        done("compliance", t0, {
            "decision": decision.decision,
            "risk_score": float(decision.risk_score),
            "risk_components": asdict(rc) if rc else {},
            "violated_rules": [_rule_to_dict(r) for r in decision.violated_rules],
            "applicable_count": len(decision.applicable_rules),
            "respected_rules": [_rule_to_dict(r) for r in decision.applicable_rules
                                 if not getattr(r, "violated", False)][:15],
            "fatf_coverage": dict(decision.fatf_coverage or {}),
            "binding_coverage": dict(decision.binding_coverage or {}),
            "human_review_flags": list(decision.human_review_flags or []),
            "symbolic_rules": int(dstats.get("symbolic_rules", 0)),
            "semantic_rules": int(dstats.get("semantic_rules", 0)),
            "rules_persisted": int(dstats.get("rules_persisted", 0)),
            # [Diagnostic + jury] gate anti-hallucination : nb de règles
            # extraites par le LLM PUIS rejetées faute d'ancrage verbatim,
            # et nb de dimensions FATF confiées au chemin sémantique.
            "semantic_rejected": int(dstats.get("semantic_rejected", 0)),
            "gap_dims": int(dstats.get("gap_dims", 0)),
        })

        # ── [6] ML (signal indicatif — divergence gérée par le SAR) ──
        t0 = time.time()
        ml_result = C["ml"].predict(tx, client_id=client_id)
        ml_public = {k: v for k, v in ml_result.items() if k != "all_features"}
        done("ml", t0, ml_public)
        # [RLHF] Persiste le vecteur de features ML (hors réponse publique) —
        # permet de transformer le feedback superviseur en exemples étiquetés
        # (rlhf.build_feedback_dataset). Aligné sur meta["features"].
        job["ml_features"] = ml_result.get("all_features", {})

        # ── [7] SAR ───────────────────────────────────────────────────
        t0 = time.time()
        sar = C["sar"].generate(
            compliance_decision=decision,
            ml_result=ml_result,
            client_result=client_result,
            transaction=tx,
            l1_signals=list(l1.risk_signals))
        sars_dir = BASE / "sars"
        sars_dir.mkdir(exist_ok=True)
        sar_md = sar.to_markdown(str(sars_dir / f"sar_{tx_id}.md"))
        sar.to_json(str(sars_dir / f"sar_{tx_id}.json"))
        done("sar", t0, {
            "sar_id": sar.sar_id,
            "risk_score_combined": float(sar.risk_score_combined),
            "risk_level": sar.risk_level,
            "recommendation": sar.recommendation,
            "divergence_flag": sar.divergence_flag,
            "divergence_action": sar.divergence_action,
            "aml_pattern": sar.aml_pattern,
            "aml_pattern_explanation": sar.aml_pattern_explanation,
            "coverage_fatf": float(sar.coverage_fatf),
            "coverage_bct": float(sar.coverage_bct),
            "required_documents": list(sar.required_documents),
        })
        job["sar_markdown"] = sar_md
        job["sar_json"] = sar.to_dict()

        # ── [8] WORKFLOW (état + file superviseur) ────────────────────
        t0 = time.time()
        outcome = C["workflow"].process({**tx}, decision, tx_id=tx_id)
        done("workflow", t0, {
            "state": outcome.state,
            "summary": outcome.summary,
            "needs_supervisor": bool(outcome.needs_supervisor),
            "str_required": bool(outcome.str_required),
            "required_documents": list(outcome.required_documents),
        })

        # ── Historique Redis (mémoire glissante) ──────────────────────
        try:
            C["history"].record(
                transaction={
                    "client_id": client_id, "sender_account_id": client_id,
                    "tx_type": tx.get("tx_type", ""), "amount": tx.get("amount", 0),
                    "currency": tx.get("currency", "TND"),
                    "destination_country": tx.get("destination_country", ""),
                    "beneficiary_name": tx.get("beneficiary_name", ""),
                    "receiver_account_id": str(
                        tx.get("beneficiary_name", "")).replace("Account-", ""),
                    "timestamp": tx.get("timestamp", ""),
                },
                decision=decision.decision,
                is_str=decision.decision in ("STR", "BLOCK"))
        except Exception as e:
            steps.setdefault("warnings", []).append(f"redis: {e}")

        # ── Résumé final du job ───────────────────────────────────────
        job["summary"] = {
            "l1_decision": l1.decision,
            "ca_decision": decision.decision,
            "ca_score": float(decision.risk_score),
            "ml_decision": ml_result.get("ml_decision", ""),
            "risk_level": sar.risk_level,
            "combined_score": float(sar.risk_score_combined),
            "state": outcome.state,
            "aml_pattern": sar.aml_pattern,
        }
        job["status"] = "DONE"
        job["finished_at"] = datetime.now().isoformat()

    except Exception as e:
        job["status"] = "ERROR"
        job["error"] = f"{e}"
        job["trace"] = traceback.format_exc()[-2000:]
    finally:
        _persist_job(job)   # l'historique survit aux redémarrages
        _publish_job_event(job, "job")   # push SSE : analyse terminée


# ═══════════════════════════════════════════════════════════════════════
#  SCHÉMAS
# ═══════════════════════════════════════════════════════════════════════

class JustifFile(BaseModel):
    filename: str
    content_b64: str


class TransactionIn(BaseModel):
    """Transaction à analyser (mêmes champs que le pipeline E2E)."""
    tx_type: str = Field("SWIFT", description="SWIFT | VIREMENT_LOCAL | MANDAT | DEPOT_ESPECES")
    amount: float = Field(..., examples=[9500])
    currency: str = Field("TND")
    sender_name: str = Field("", description="ex: Account-802559360 (le client_id en est extrait)")
    sender_country: str = Field("Tunisia")
    beneficiary_name: str = Field("")
    destination_country: str = Field("Tunisia")
    motive: str = Field("", description="motif de l'opération")
    has_justification: Optional[bool] = None
    timestamp: Optional[str] = None
    client_id: Optional[str] = Field(None, description="prioritaire sur sender_name pour le profil")
    sender_rib: str = Field("", description="RIB/IBAN émetteur (validé côté front)")
    beneficiary_iban: str = Field("", description="RIB/IBAN bénéficiaire")
    # Justificatifs joints À LA SOUMISSION (facture + contrat…) — analysés en
    # étape 0 du pipeline ; leur verdict REMPLACE la case déclarative
    # has_justification (le justificatif n'est plus déclaré, il est VÉRIFIÉ).
    justificatifs: Optional[list[JustifFile]] = None


class ValidateIn(BaseModel):
    action: str = Field(..., description="APPROVE | REJECT | ANNOTATE")
    label: str = Field("", description="FALSE_POSITIVE | TRUE_FRAUD | CORRECT | ...")
    annotation: str = Field("", description="justification (obligatoire côté front)")
    supervisor_id: str = Field("supervisor")


# ═══════════════════════════════════════════════════════════════════════
#  APPLICATION
# ═══════════════════════════════════════════════════════════════════════

app = FastAPI(
    title="AML/KYC Pipeline API",
    version="1.0",
    description=("Backend du système AML/KYC Agentic RAG — expose les 8 étapes "
                 "réelles du pipeline (L1 → Client → Planning → Retrieval → "
                 "Compliance → ML → SAR → Workflow) pour le frontend React."))

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://localhost:3000",
                   "http://127.0.0.1:5173"],
    allow_methods=["*"], allow_headers=["*"])

app.include_router(auth_router)   # /api/auth/register · verify · login · me


def _job_retraining_cycle() -> dict:
    """Cycle hebdomadaire d'amélioration continue (feedback superviseur)."""
    from retraining import run_cycle
    report = run_cycle()
    return {
        "labelled": report["stats"]["labelled"],
        "agreement_rate": report["stats"]["agreement_rate"],
        "false_positive_rate": report["stats"]["false_positive_rate"],
        "retrain": report["decision"]["retrain"],
        "reason": report["decision"]["reason"],
        "training_status": report["training"]["status"],
    }


def _job_drift_check() -> dict:
    """Contrôle quotidien : ne réentraîne que si une dérive est avérée.

    C'est le garde-fou qui évite d'attendre le cycle hebdomadaire quand le
    taux de fausses alertes s'envole.
    """
    from retraining import load_feedback, enrich_with_transactions, measure, FP_ALERT
    stats = measure(enrich_with_transactions(load_feedback()))
    drift = stats["labelled"] > 0 and stats["false_positive_rate"] > FP_ALERT
    if drift:
        print(f"  [SCHED] dérive détectée "
              f"({stats['false_positive_rate']:.1%} de fausses alertes) "
              f"→ cycle de réentraînement anticipé")
        return {"drift": True, **_job_retraining_cycle()}
    return {"drift": False, **stats}


def _job_weekly_report() -> dict:
    """Rapport hebdomadaire envoyé automatiquement au responsable conformité."""
    stats = _weekly_stats(7)
    dest = os.getenv("WEEKLY_REPORT_TO", "").strip()
    if not dest:
        dest = _default_report_recipient()
    if not dest:
        return {"sent": False, "reason": "aucun destinataire déterminé"}
    send_weekly_report_email(dest, stats)
    return {"sent": True, "recipient": dest, "total": stats.get("total")}


def _default_report_recipient() -> str:
    """Premier Compliance Officer (à défaut, admin) enregistré."""
    try:
        from api.auth import _db, _db_lock
        with _db_lock, _db() as conn:
            row = conn.execute(
                "SELECT email FROM users WHERE role='COMPLIANCE_OFFICER' "
                "AND email IS NOT NULL ORDER BY rowid LIMIT 1").fetchone()
            if not row:
                row = conn.execute(
                    "SELECT email FROM users WHERE role='ADMIN' "
                    "AND email IS NOT NULL ORDER BY rowid LIMIT 1").fetchone()
        return row["email"] if row else ""
    except Exception:
        return ""


@app.on_event("startup")
def _startup():
    auth_init_db()   # table users (SQLite platform.db)
    _load_jobs()     # historique des dossiers (survit aux redémarrages)

    # ── Tâches périodiques (Tâche 1) ────────────────────────────────
    #  Cadence RETENUE : hebdomadaire pour le réentraînement — sur 3 jours le
    #  volume d'étiquettes humaines est trop faible et le modèle apprendrait
    #  du bruit. Un contrôle quotidien permet un déclenchement anticipé si le
    #  taux de fausses alertes dérive.
    SCHEDULER.register(Job(
        name="retraining", func=_job_retraining_cycle,
        weekday=SUNDAY, hour=2, label="Cycle de réentraînement (RLHF)"))
    SCHEDULER.register(Job(
        name="drift_check", func=_job_drift_check,
        weekday=None, hour=3, label="Contrôle de dérive (quotidien)"))
    SCHEDULER.register(Job(
        name="weekly_report", func=_job_weekly_report,
        weekday=FRIDAY, hour=17, label="Rapport hebdomadaire par email"))
    SCHEDULER.start()

    # Init en arrière-plan pour ne pas bloquer le démarrage du serveur
    threading.Thread(target=init_components, daemon=True).start()


@app.get("/api/scheduler/status")
def scheduler_status(authorization: Optional[str] = Header(None)):
    """État du planificateur : tâches, dernières exécutions, historique."""
    user = current_user(authorization)
    if not user or user["r"] not in ("COMPLIANCE_OFFICER", "ADMIN"):
        raise HTTPException(403, "Accès réservé à la conformité et à l'audit.")
    return SCHEDULER.status()


@app.post("/api/scheduler/run/{job_name}")
def scheduler_run(job_name: str, authorization: Optional[str] = Header(None)):
    """Déclenche manuellement une tâche planifiée (démonstration au jury)."""
    user = current_user(authorization)
    if not user or user["r"] not in ("COMPLIANCE_OFFICER", "ADMIN"):
        raise HTTPException(403, "Accès réservé à la conformité.")
    return SCHEDULER.run_job(job_name, manual=True)


# ── RLHF — boucle d'amélioration continue (feedback superviseur) ─────────
@app.get("/api/rlhf/summary")
def rlhf_summary(authorization: Optional[str] = Header(None)):
    """État de la boucle RLHF : accord système/expert, taux de fausses alertes,
    métriques du modèle courant, dernier cycle."""
    user = current_user(authorization)
    if not user or user["r"] not in ("COMPLIANCE_OFFICER", "ADMIN"):
        raise HTTPException(403, "Accès réservé à la conformité et à l'audit.")
    from retraining import (load_feedback, enrich_with_transactions, measure,
                            load_state, MIN_LABELS, FP_ALERT)
    stats = measure(enrich_with_transactions(load_feedback()))
    meta = {}
    try:
        meta = json.loads((BASE / "aml_ml_demo_meta.json").read_text(encoding="utf-8"))
    except Exception:
        pass
    return {
        "stats": stats,
        "model": {"threshold": meta.get("threshold"),
                  "metrics": meta.get("metrics", {}),
                  "last_rlhf": meta.get("last_rlhf")},
        "thresholds": {"min_labels": MIN_LABELS, "fp_alert": FP_ALERT},
        "last_cycle": load_state(),
    }


@app.post("/api/rlhf/run")
def rlhf_run(force: bool = False, authorization: Optional[str] = Header(None)):
    """Lance un cycle RLHF complet (agrégation → décision → warm-start garde-fou
    → trace). `force=true` bypasse le seuil de volume pour la démonstration."""
    user = current_user(authorization)
    if not user or user["r"] not in ("COMPLIANCE_OFFICER", "ADMIN"):
        raise HTTPException(403, "Accès réservé à la conformité.")
    from retraining import run_cycle
    return run_cycle(force=force)


@app.get("/api/health")
def health():
    return {"status": "ok" if C.get("ready") else "initializing",
            "components_ready": bool(C.get("ready")),
            "neo4j": bool(C.get("neo4j")),
            "jobs": len(JOBS)}


@app.post("/api/analyze")
def analyze(tx: TransactionIn, authorization: Optional[str] = Header(None)):
    """Lance l'analyse complète (asynchrone). Le front sonde GET /api/result/{tx_id}.

    Si un token est fourni (Authorization: Bearer …), la transaction est
    rattachée à son propriétaire → l'espace utilisateur ne voit que SES dossiers.
    """
    if not C.get("ready"):
        raise HTTPException(503, "Composants en cours d'initialisation — réessayer dans quelques secondes")
    user = current_user(authorization)
    tx_id = _next_tx_id()
    payload = tx.model_dump()
    # Les fichiers (base64, volumineux) sont gardés HORS de transaction/steps
    files = payload.pop("justificatifs", None) or []
    JOBS[tx_id] = {
        "tx_id": tx_id, "status": "RUNNING", "current_step": None,
        "transaction": payload, "steps": {},
        "_files": files,
        "source": "USER",
        "owner": user["u"] if user else None,
        "created_at": datetime.now().isoformat(),
    }
    threading.Thread(target=execute_pipeline,
                     args=(payload, tx_id), daemon=True).start()
    return {"tx_id": tx_id, "status": "RUNNING"}


@app.get("/api/result/{tx_id}")
def result(tx_id: str):
    """Statut + sorties de chaque étape (le front peut afficher au fil de l'eau)."""
    job = JOBS.get(tx_id)
    if not job:
        raise HTTPException(404, f"Transaction {tx_id} inconnue")
    return {k: v for k, v in job.items()
            if k not in ("sar_markdown", "_files", "ml_features")}


@app.get("/api/sar/{tx_id}/markdown", response_class=PlainTextResponse)
def sar_markdown(tx_id: str):
    job = JOBS.get(tx_id)
    if not job or "sar_markdown" not in job:
        raise HTTPException(404, "SAR non disponible")
    return job["sar_markdown"]


@app.get("/api/transactions")
def transactions(mine: bool = False, source: Optional[str] = None,
                 authorization: Optional[str] = Header(None)):
    """Historique des analyses (résumés, plus récentes d'abord).

    Filtres : ?mine=true (avec token → uniquement MES dossiers, espace
    utilisateur) · ?source=SIMULATOR|USER (espace admin).
    """
    user = current_user(authorization)
    out = []
    for job in JOBS.values():
        if mine and (not user or job.get("owner") != user["u"]):
            continue
        if source and job.get("source") != source:
            continue
        out.append(_tx_item(job))
    return sorted(out, key=lambda j: j.get("created_at") or "", reverse=True)


def _tx_item(job: dict) -> dict:
    """Projection d'un dossier vers l'item de liste du frontend (TxListItem)."""
    tx = job.get("transaction", {})
    return {
        "tx_id": job.get("tx_id"), "status": job.get("status"),
        "created_at": job.get("created_at"),
        "source": job.get("source", "USER"),
        "owner": job.get("owner"),
        "transaction": {k: tx.get(k) for k in
                        ("tx_type", "amount", "currency",
                         "sender_country", "destination_country")},
        "summary": job.get("summary"),
        "supervisor_decision": job.get("supervisor_decision"),
    }


@app.get("/api/transactions/query")
def transactions_query(
    min_amount: Optional[float] = None,
    decision: Optional[str] = None,
    client: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    limit: int = 300,
    authorization: Optional[str] = Header(None),
):
    """Requête COMPLEXE sur les dossiers — démonstration PostgreSQL (Tâche 4).

    Filtre par montant minimum, décision, client (owner), et période — traduit
    en SQL indexé côté PostgreSQL (repli filtrage mémoire si base absente).
    Réservé à l'audit / la conformité (séparation des devoirs).
    """
    user = current_user(authorization)
    if not user or user["r"] not in ("COMPLIANCE_OFFICER", "ADMIN"):
        raise HTTPException(403, "Accès réservé à l'audit et à la conformité.")
    # date_to au format jour seul (YYYY-MM-DD) → inclure toute la journée
    if date_to and len(date_to) == 10:
        date_to = f"{date_to}T23:59:59"
    jobs = STORE.query(
        owner=(client or None), decision=(decision or None),
        min_amount=min_amount, date_from=(date_from or None),
        date_to=(date_to or None), limit=limit)
    return {
        "count": len(jobs),
        "backend": STORE.kind,             # « PostgreSQL » ou « fichiers JSON »
        "results": [_tx_item(j) for j in jobs],
    }


@app.post("/api/validate/{tx_id}")
def validate(tx_id: str, body: ValidateIn,
             authorization: Optional[str] = Header(None)):
    """Décision superviseur → feedback RLHF (output/supervisor_feedback.jsonl).

    SÉPARATION DES DEVOIRS (segregation of duties) : seul le rôle
    COMPLIANCE_OFFICER peut approuver/refuser un dossier — un ADMIN (profil
    technique DSI) ne doit pas pouvoir valider une décision réglementaire.
    """
    user = current_user(authorization)
    if not user:
        raise HTTPException(401, "Connexion requise pour valider un dossier")
    if user["r"] != "COMPLIANCE_OFFICER":
        raise HTTPException(403, "Seul un Compliance Officer peut valider un "
                                 "dossier (séparation des devoirs) — rôle "
                                 f"actuel : {user['r']}")
    job = JOBS.get(tx_id)
    if not job:
        raise HTTPException(404, f"Transaction {tx_id} inconnue")
    evt = C["workflow"].queue.record_decision(
        tx_id, supervisor_action=body.action, supervisor_label=body.label,
        annotation=body.annotation, supervisor_id=body.supervisor_id)
    job["supervisor_decision"] = {
        "action": body.action, "label": body.label,
        "annotation": body.annotation, "by": body.supervisor_id,
        "at": evt.created_at,
    }
    # ÉTAT DU DOSSIER : la décision superviseur fait AVANCER la machine à
    # états (sinon l'utilisateur voyait toujours « en attente » côté espace
    # utilisateur alors que le dossier était déjà tranché).
    new_state = {"APPROVE": "APPROVED_BY_SUPERVISOR",
                 "REJECT": "REJECTED_BY_SUPERVISOR"}.get(body.action)
    if new_state:
        job.setdefault("summary", {})["state"] = new_state
        wf = job.get("steps", {}).get("workflow")
        if isinstance(wf, dict):
            wf["state"] = new_state
            wf["summary"] = ("Dossier APPROUVÉ par le superviseur — l'opération "
                             "peut être exécutée."
                             if body.action == "APPROVE" else
                             "Dossier REFUSÉ par le superviseur — informer le client.")
    _persist_job(job)
    # EMAIL DE DÉCISION au propriétaire du dossier (HTML vert/rouge) —
    # en thread pour ne pas ralentir la réponse ; jamais bloquant.
    # L'owner est le USERNAME du dossier (parfois « user1 », pas un email) →
    # on résout l'adresse réelle avant l'envoi (sinon aucune notification).
    owner = job.get("owner")
    recipient = get_user_email(owner) if owner else None
    if new_state and recipient:
        threading.Thread(
            target=send_decision_email,
            args=(recipient, tx_id, body.action == "APPROVE",
                  job.get("transaction", {}), body.annotation),
            daemon=True).start()
        print(f"  [API] 📧 email de décision → {recipient} "
              f"({'APPROUVÉ' if body.action == 'APPROVE' else 'REFUSÉ'})")
    elif new_state:
        print(f"  [API] ⚠️ email de décision NON envoyé : owner={owner!r} "
              f"sans adresse email résolue")
    _publish_job_event(job, "decision")   # push SSE : décision superviseur
    return {"tx_id": tx_id, "recorded": True, "feedback": asdict(evt)}


# ── FILE SUPERVISEUR : priorités + tri par urgence + historique (Tâche 3) ───
#  La file est DÉRIVÉE des dossiers persistés (PostgreSQL) → elle survit aux
#  redémarrages. Priorité selon la décision système ; tri = priorité puis
#  ancienneté (le plus ancien/urgent d'abord).
_PRIORITY = {"STR": ("HAUTE", 0), "BLOCK": ("HAUTE", 0), "EDD": ("MOYENNE", 1)}


def _priority_of(decision: Optional[str]) -> tuple:
    return _PRIORITY.get(decision or "", ("NORMALE", 2))


def _age_hours(created) -> Optional[float]:
    """Heures écoulées depuis la création du dossier (pour le tri d'urgence)."""
    if not created:
        return None
    try:
        dt = datetime.fromisoformat(str(created).replace("Z", "+00:00"))
    except ValueError:
        return None
    if dt.tzinfo:
        dt = dt.replace(tzinfo=None)
    return round((datetime.now() - dt).total_seconds() / 3600, 1)


@app.get("/api/supervisor/queue")
def supervisor_queue(authorization: Optional[str] = Header(None)):
    """File d'attente enrichie : priorité (HAUTE/MOYENNE/NORMALE) + ancienneté,
    triée par urgence. Réservée à la conformité / l'audit."""
    user = current_user(authorization)
    if not user or user["r"] not in ("COMPLIANCE_OFFICER", "ADMIN"):
        raise HTTPException(403, "Accès réservé à la conformité et à l'audit.")
    items = []
    for job in JOBS.values():
        if job.get("status") != "DONE":
            continue
        state = (job.get("summary") or {}).get("state") or ""
        if not state.startswith("PENDING") or job.get("supervisor_decision"):
            continue
        decision = (job.get("summary") or {}).get("ca_decision")
        label, rank = _priority_of(decision)
        item = _tx_item(job)
        item["priority"] = label
        item["priority_rank"] = rank
        item["age_hours"] = _age_hours(job.get("created_at"))
        items.append(item)
    # tri : priorité (HAUTE d'abord) puis ancienneté (plus ancien d'abord)
    items.sort(key=lambda x: (x["priority_rank"], x.get("created_at") or ""))
    return items


@app.get("/api/supervisor/history")
def supervisor_history(limit: int = 100,
                       authorization: Optional[str] = Header(None)):
    """Historique des décisions superviseur passées (plus récentes d'abord)."""
    user = current_user(authorization)
    if not user or user["r"] not in ("COMPLIANCE_OFFICER", "ADMIN"):
        raise HTTPException(403, "Accès réservé à la conformité et à l'audit.")
    out = [_tx_item(job) for job in JOBS.values()
           if job.get("supervisor_decision")]
    out.sort(key=lambda x: (x.get("supervisor_decision") or {}).get("at") or "",
             reverse=True)
    return out[:limit]


@app.get("/api/events")
async def events(token: Optional[str] = None):
    """Flux SSE temps réel (Tâche 2) — pousse une notification à chaque
    changement d'état d'un dossier (fin d'analyse, décision superviseur),
    remplaçant le polling. Le token passe en query param car EventSource
    n'autorise pas d'en-tête Authorization."""
    user = current_user(f"Bearer {token}") if token else None
    if not user:
        raise HTTPException(401, "Token requis pour le flux d'événements.")
    q = await BUS.subscribe()

    async def gen():
        try:
            yield "event: ready\ndata: {}\n\n"        # confirme l'ouverture
            while True:
                try:
                    event = await asyncio.wait_for(q.get(), timeout=25)
                    yield f"data: {json.dumps(event, ensure_ascii=False, default=str)}\n\n"
                except asyncio.TimeoutError:
                    yield ": keep-alive\n\n"           # évite la coupure proxy
        finally:
            BUS.unsubscribe(q)

    return StreamingResponse(gen(), media_type="text/event-stream", headers={
        "Cache-Control": "no-cache",
        "Connection": "keep-alive",
        "X-Accel-Buffering": "no",
    })


def _weekly_stats(days: int = 7) -> dict:
    """Agrège les analyses des `days` derniers jours pour le rapport hebdo."""
    from datetime import timedelta
    since = datetime.now() - timedelta(days=days)
    by_dec, by_pat = {}, {}
    total = legit = flagged = blocked = pending = 0
    for job in JOBS.values():
        if job.get("status") != "DONE":
            continue
        try:
            if datetime.fromisoformat(job["created_at"]) < since:
                continue
        except Exception:
            continue
        s = job.get("summary", {}) or {}
        dec = s.get("ca_decision", "?")
        by_dec[dec] = by_dec.get(dec, 0) + 1
        by_pat[s.get("aml_pattern", "?")] = by_pat.get(s.get("aml_pattern", "?"), 0) + 1
        total += 1
        if dec in ("PASS", "REVIEW", "KYC", "RECORD"):
            legit += 1
        if dec in ("EDD", "STR"):
            flagged += 1
        if dec == "BLOCK":
            blocked += 1
        if str(s.get("state", "")).startswith("PENDING"):
            pending += 1
    top_pattern = max((p for p in by_pat.items() if p[0] not in ("NORMAL", "?")),
                      key=lambda x: x[1], default=("aucun", 0))[0]
    return {"total": total, "legit": legit, "flagged": flagged, "blocked": blocked,
            "pending": pending, "by_decision": by_dec, "top_pattern": top_pattern,
            "period": f"7 derniers jours (au {datetime.now().strftime('%d/%m/%Y')})"}


class WeeklyReportIn(BaseModel):
    email: str = ""   # destinataire ; défaut = l'utilisateur connecté


@app.post("/api/reports/weekly")
def weekly_report(body: WeeklyReportIn,
                  authorization: Optional[str] = Header(None)):
    """Envoie le rapport hebdomadaire par email (Compliance Officer / Admin).

    En production : un job planifié (APScheduler / cron) appelle cet endpoint
    chaque vendredi. Ici, déclenchement manuel depuis l'interface + aperçu JSON."""
    user = current_user(authorization)
    if not user or user["r"] not in ("COMPLIANCE_OFFICER", "ADMIN"):
        raise HTTPException(403, "Réservé au Compliance Officer / Admin")
    dest = (body.email or "").strip() or user["u"]
    st = _weekly_stats()
    sent = False
    if "@" in dest:
        threading.Thread(target=send_weekly_report_email,
                         args=(dest, st), daemon=True).start()
        sent = True
    return {"sent": sent, "recipient": dest, "stats": st}


@app.get("/api/stats")
def stats():
    """KPIs pour le dashboard (agrégés sur les analyses de la session)."""
    done = [j for j in JOBS.values() if j.get("status") == "DONE"]
    by_decision, by_level, by_pattern = {}, {}, {}
    total_ms = 0
    for j in done:
        s = j.get("summary", {})
        by_decision[s.get("ca_decision", "?")] = by_decision.get(s.get("ca_decision", "?"), 0) + 1
        by_level[s.get("risk_level", "?")] = by_level.get(s.get("risk_level", "?"), 0) + 1
        by_pattern[s.get("aml_pattern", "?")] = by_pattern.get(s.get("aml_pattern", "?"), 0) + 1
        total_ms += sum(st.get("duration_ms", 0)
                        for st in (j.get("steps") or {}).values()
                        if isinstance(st, dict))
    pending = [j["tx_id"] for j in done
               if j.get("summary", {}).get("state", "").startswith("PENDING")
               and not j.get("supervisor_decision")]
    return {
        "total_analyzed": len(done),
        "running": sum(1 for j in JOBS.values() if j.get("status") == "RUNNING"),
        "errors": sum(1 for j in JOBS.values() if j.get("status") == "ERROR"),
        "by_decision": by_decision,
        "by_risk_level": by_level,
        "by_pattern": by_pattern,
        "pending_supervisor": pending,
        "avg_pipeline_ms": int(total_ms / max(len(done), 1)),
    }


@app.get("/api/clients/{client_id}")
def client(client_id: str):
    """Profil KYC (pré-remplissage du formulaire) + historique comportemental."""
    if not C.get("ready"):
        raise HTTPException(503, "Initialisation en cours")
    if not C["clients_db"].exists(client_id):
        # Les nouveaux clients (page « Mon profil ») sont upsertés dans
        # clients_db.json par api/auth.py → recharger le cache du loader.
        try:
            C["clients_db"]._load()
        except Exception:
            pass
    known = C["clients_db"].exists(client_id)
    profile = C["clients_db"].get_profile(client_id, fallback_default=True)
    if profile is None:
        raise HTTPException(404, "Client inconnu")
    prof_hist = C["history"].behavioral_profile(client_id)
    return {
        "client_id": client_id, "known": bool(known),
        "nationality": getattr(profile, "nationality", ""),
        "country_residence": getattr(profile, "country_residence", ""),
        "is_pep": bool(getattr(profile, "is_pep", False)),
        "pep_type": getattr(profile, "pep_type", None),
        "kyc_complete": bool(getattr(profile, "kyc_complete", False)),
        "previous_strs": getattr(profile, "previous_strs", 0),
        "account_type": getattr(profile, "account_type", ""),
        "history": prof_hist.to_history_dict(),
    }


# ═══════════════════════════════════════════════════════════════════════
#  JUSTIFICATIF — upload (base64, zéro dépendance) + analyse de conformité
# ═══════════════════════════════════════════════════════════════════════
#  Méthode HYBRIDE (plus robuste qu'un LLM seul) :
#   1. EXTRACTION   : docling si installé (meilleur sur les tableaux),
#                     sinon pypdf (PDF) / python-docx (DOCX) / décodage (TXT).
#   2. DÉTERMINISTE : montant, noms émetteur/bénéficiaire, date récente
#                     cherchés par normalisation/regex → auditable, zéro LLM.
#   3. LLM SOUS CONTRAINTE : juge la COHÉRENCE sémantique (type de document,
#                     correspondance avec l'opération et les documents requis)
#                     à partir du SEUL texte extrait. Repli déterministe.

class JustificatifIn(BaseModel):
    filename: str = Field(..., description="nom du fichier (.pdf .txt .docx)")
    content_b64: str = Field(..., description="contenu encodé base64")


_DOCLING_CONVERTER = None


def _docling_converter():
    """Convertisseur docling en singleton, OCR désactivé.

    Nos justificatifs sont TOUJOURS des documents nés numériques (factures,
    contrats, attestations générés/exportés) — jamais des scans papier.
    `do_ocr` vaut True par DÉFAUT dans docling : il lance alors un moteur OCR
    (RapidOCR) sur CHAQUE page, y compris celles qui ont déjà du texte
    exploitable → télécharge ~40 Mo de modèles au 1er appel et ajoute
    plusieurs minutes de traitement à chaque upload. On le désactive tout en
    gardant `do_table_structure=True` (la détection des tableaux, essentielle
    pour les factures/contrats).
    """
    global _DOCLING_CONVERTER
    if _DOCLING_CONVERTER is None:
        from docling.document_converter import DocumentConverter, PdfFormatOption
        from docling.datamodel.pipeline_options import PdfPipelineOptions
        from docling.datamodel.base_models import InputFormat
        opts = PdfPipelineOptions(do_ocr=False, do_table_structure=True)
        _DOCLING_CONVERTER = DocumentConverter(
            format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=opts)})
    return _DOCLING_CONVERTER


PDF_DOCLING_TIMEOUT = 90   # s — au-delà, repli pypdf (l'analyse ne bloque jamais)


def _extract_doc_text(filename: str, data: bytes) -> str:
    """Texte brut du document — DOCLING D'ABORD, replis garantis (mesuré) :

      .pdf  → DOCLING (layout + TableFormer : tableaux → markdown structuré ;
              4,3 s à froid / 1,0 s à chaud une fois le modèle en cache).
              Repli PYPDF si échec OU si > PDF_DOCLING_TIMEOUT (ex. modèle
              non téléchargé : le 1er téléchargement avait pris 47 min).
      .docx → DOCLING (0,1 s, tableaux Word natifs), repli python-docx
              (qui lit aussi les cellules de tableaux).
      .txt  → décodage direct.
    """
    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        # 1) docling — indispensable pour les factures/contrats à TABLEAUX
        #    (pypdf aplatit les tableaux en texte désordonné)
        try:
            import tempfile
            from concurrent.futures import ThreadPoolExecutor
            from concurrent.futures import TimeoutError as _FutTimeout
            with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as f:
                f.write(data)
                tmp = f.name
            pool = ThreadPoolExecutor(max_workers=1)
            try:
                fut = pool.submit(
                    lambda: _docling_converter().convert(tmp)
                    .document.export_to_markdown())
                return fut.result(timeout=PDF_DOCLING_TIMEOUT)
            finally:
                pool.shutdown(wait=False)
                try:
                    Path(tmp).unlink(missing_ok=True)
                except Exception:
                    pass   # fichier encore ouvert par le thread → nettoyé par l'OS
        except ImportError:
            pass
        except _FutTimeout:
            print(f"  [API] ⚠️ docling > {PDF_DOCLING_TIMEOUT}s sur {filename} "
                  f"(modèle en cours de téléchargement ?) — repli pypdf")
        except Exception as e:
            print(f"  [API] ⚠️ docling a échoué sur {filename} ({str(e)[:80]}) — repli pypdf")
        # 2) repli pypdf — texte brut (toujours disponible, <0,5 s)
        from io import BytesIO
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages[:15])

    if ext == ".docx":
        # 1) docling — meilleur parsing structurel (tableaux → markdown)
        try:
            import tempfile
            with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as f:
                f.write(data)
                tmp = f.name
            try:
                return _docling_converter().convert(tmp).document.export_to_markdown()
            finally:
                Path(tmp).unlink(missing_ok=True)
        except ImportError:
            pass
        except Exception as e:
            print(f"  [API] ⚠️ docling a échoué sur {filename} ({str(e)[:80]}) — repli python-docx")
        # 2) repli python-docx (paragraphes + cellules de tableaux)
        try:
            from io import BytesIO
            import docx
            d = docx.Document(BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for table in d.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)
        except ImportError:
            raise HTTPException(400, "Format .docx : installer python-docx "
                                     "(pip install python-docx) ou fournir un PDF")

    # .txt / .md / autre texte
    for enc in ("utf-8", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise HTTPException(400, f"Format {ext or 'inconnu'} non pris en charge "
                             "(PDF, DOCX ou TXT)")


def _amounts_in_text(text: str) -> set:
    """Tous les MONTANTS du document, parsés en valeurs numériques exactes.

    Remplace le test par sous-chaîne qui produisait des faux positifs :
    « 1200 » était trouvé DANS « 12 000,00 » → une facture de 12 000 USD
    validait une transaction de 1 200 USD (constaté en démo). Ici chaque
    token numérique est parsé (séparateurs de milliers espace/point/virgule,
    décimales 1-2 chiffres) puis comparé en valeur.
    """
    import re as _re
    out = set()
    for m in _re.finditer(r"\d[\d\s.,']{0,15}\d|\d", text):
        tok = _re.sub(r"[\s']", "", m.group(0))
        # dernier '.' ou ',' suivi de 1-2 chiffres en fin = séparateur décimal
        dec = _re.match(r"^(\d[\d.,]*?)[.,](\d{1,2})$", tok)
        if dec:
            try:
                out.add(float(f"{_re.sub(r'[.,]', '', dec.group(1))}.{dec.group(2)}"))
            except ValueError:
                pass
        try:
            out.add(float(_re.sub(r"[.,]", "", tok)))
        except ValueError:
            pass
    return out


def _deterministic_checks(text: str, tx: dict) -> dict:
    """Vérifications auditables sans LLM : montant / noms / date récente."""
    import re as _re
    from datetime import timedelta
    low = text.lower()
    digits = _re.sub(r"[  ,.'']", "", text)

    amount = float(tx.get("amount") or 0)
    # Comparaison en VALEUR exacte (fix du faux positif par sous-chaîne)
    amount_found = bool(amount) and any(
        abs(v - amount) < 0.01 for v in _amounts_in_text(text))

    def name_found(name):
        toks = [t for t in _re.split(r"[\s\-]+", str(name or "").lower())
                if len(t) >= 3 and t != "account"]
        return bool(toks) and all(t in low for t in toks)

    dates = []
    for m in _re.finditer(
            r"\b(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})\b|\b(\d{4})-(\d{2})-(\d{2})\b",
            text):
        try:
            if m.group(3):
                dates.append(datetime(int(m.group(3)), int(m.group(2)), int(m.group(1))))
            else:
                dates.append(datetime(int(m.group(4)), int(m.group(5)), int(m.group(6))))
        except ValueError:
            continue
    now = datetime.now()
    date_recente = any(abs((now - d).days) <= 365 for d in dates)

    return {
        "montant_present": amount_found,
        "emetteur_present": name_found(tx.get("sender_name")),
        "beneficiaire_present": name_found(tx.get("beneficiary_name")),
        "date_recente_presente": date_recente,
        "dates_trouvees": len(dates),
    }


def _llm_review_justificatif(text: str, tx: dict, required_docs: list,
                             checks: dict) -> Optional[dict]:
    """Jugement sémantique LLM (sous contrainte) — None si indisponible."""
    try:
        from llm_client import get_llm
        llm = get_llm()
    except Exception:
        return None
    if llm is None:
        return None
    prompt = (
        "Tu es analyste conformité bancaire. Un client a fourni un document "
        "justificatif pour la transaction ci-dessous. Évalue si ce document "
        "est un justificatif économique RECEVABLE.\n"
        "CONTRAINTES : base-toi UNIQUEMENT sur le texte extrait fourni. "
        "N'invente rien. Réponds en français.\n\n"
        f"## Transaction\n{json.dumps({k: tx.get(k) for k in ('tx_type', 'amount', 'currency', 'sender_name', 'beneficiary_name', 'destination_country', 'motive')}, ensure_ascii=False)}\n\n"
        f"## Documents demandés par la conformité\n{json.dumps(required_docs, ensure_ascii=False)}\n\n"
        f"## Vérifications automatiques déjà faites\n{json.dumps(checks, ensure_ascii=False)}\n\n"
        f"## Texte extrait du document (tronqué)\n{text[:3500]}\n\n"
        '## Sortie (JSON strict, sans markdown) : {"document_type": "facture|contrat|attestation|releve|autre", '
        '"conforme": true/false, "verifications": [{"critere": "...", "statut": "OK|MANQUE|DOUTE", "commentaire": "..."}], '
        '"resume": "2 phrases max : le document justifie-t-il cette opération et pourquoi"}'
    )
    try:
        import re as _re
        resp = llm.chat.completions.create(
            model=None, messages=[{"role": "user", "content": prompt}],
            temperature=0.0, max_tokens=500)
        raw = _re.sub(r"```json\s*|```", "",
                      (resp.choices[0].message.content or "").strip())
        out = json.loads(raw)
        if isinstance(out, dict) and "conforme" in out:
            return out
    except Exception as e:
        print(f"  [API] analyse LLM justificatif indisponible : {str(e)[:80]}")
    return None


def _analyze_one_justificatif(filename: str, content_b64: str, tx: dict,
                              required: list) -> dict:
    """Analyse UN document (extraction + vérifs déterministes + LLM).
    Lève ValueError avec message lisible si le fichier est inexploitable."""
    import base64
    try:
        data = base64.b64decode(content_b64.split(",")[-1])
    except Exception:
        raise ValueError("Contenu base64 invalide")
    if len(data) > 8 * 1024 * 1024:
        raise ValueError("Fichier trop volumineux (max 8 Mo)")
    text = _extract_doc_text(filename, data)
    if len(text.strip()) < 30:
        raise ValueError("Impossible d'extraire du texte exploitable "
                         "(document scanné/image ? fournir un PDF texte)")
    checks = _deterministic_checks(text, tx)
    llm_review = _llm_review_justificatif(text, tx, required, checks)
    # VERDICT DURCI : le montant EXACT doit figurer dans le document ET le
    # document doit citer une des PARTIES (émetteur OU bénéficiaire). La date
    # récente seule ne prouve pas l'appartenance — sans ce durcissement, la
    # facture d'un tiers au mauvais montant passait (constaté en démo).
    parties_ok = checks["emetteur_present"] or checks["beneficiaire_present"]
    conforme_base = checks["montant_present"] and parties_ok
    if llm_review is not None:
        conforme = bool(llm_review.get("conforme")) and conforme_base
    else:
        conforme = conforme_base
    return {
        "filename": filename,
        "size_kb": round(len(data) / 1024, 1),
        "uploaded_at": datetime.now().isoformat(),
        "extracted_chars": len(text),
        "checks": checks,
        "llm": llm_review,                      # null si LLM indisponible
        "conforme": bool(conforme),
        "new_state": None,
    }


def _maybe_unblock(job: dict, result: dict):
    """Un document recevable débloque PENDING_DOCUMENTS → PENDING_REVIEW."""
    if result["conforme"] and job.get("summary", {}).get("state") == "PENDING_DOCUMENTS":
        job["summary"]["state"] = "PENDING_REVIEW"
        if "workflow" in job.get("steps", {}):
            job["steps"]["workflow"]["state"] = "PENDING_REVIEW"
            job["steps"]["workflow"]["summary"] = (
                "Justificatif reçu et jugé recevable — dossier transmis à la "
                "revue superviseur.")
        result["new_state"] = "PENDING_REVIEW"


@app.post("/api/justificatif/{tx_id}")
def upload_justificatif(tx_id: str, body: JustificatifIn,
                        authorization: Optional[str] = Header(None)):
    """Reçoit UN justificatif (base64), l'analyse et l'AJOUTE au dossier
    (upload multiple = appels successifs : facture, puis contrat…).
    Un document recevable débloque le dossier (PENDING_DOCUMENTS → PENDING_REVIEW)."""
    job = JOBS.get(tx_id)
    if not job:
        raise HTTPException(404, f"Transaction {tx_id} inconnue")
    user = current_user(authorization)
    if job.get("owner") and (not user or user["u"] != job["owner"]):
        raise HTTPException(403, "Ce dossier appartient à un autre utilisateur")

    tx = job.get("transaction", {})
    wf = job.get("steps", {}).get("workflow", {}) or {}
    required = [str(d) for d in wf.get("required_documents", [])]
    try:
        result = _analyze_one_justificatif(body.filename, body.content_b64,
                                           tx, required)
    except ValueError as e:
        raise HTTPException(422, str(e))

    _maybe_unblock(job, result)
    job.setdefault("justificatifs", []).append(result)
    job["justificatif"] = result   # dernier document (compat front)
    _persist_job(job)
    return result


@app.get("/api/fatf-lists")
def fatf_lists():
    """Listes FATF réelles (data/fatf_lists.json) — alertes métier du
    formulaire de soumission (liste noire = rouge, grise = orange)."""
    path = BASE / "data" / "fatf_lists.json"
    if not path.exists():
        return {"black_names": {}, "grey_names": {}}
    d = json.loads(path.read_text(encoding="utf-8"))
    return {"black_names": d.get("black_names", {}),
            "grey_names": d.get("grey_names", {}),
            "last_updated": d.get("last_updated", "")}


@app.get("/api/demo-transactions")
def demo_transactions():
    """Les 6 transactions de test LI (pour le mode simulateur du front)."""
    path = BASE / "test_transactions_e2e.json"
    if not path.exists():
        return []
    return json.loads(path.read_text(encoding="utf-8"))


@app.get("/api/regulations/search")
def regulations_search(q: str, top_k: int = 5):
    """Explorateur réglementaire : recherche hybride dans le corpus (6 678 chunks)."""
    if not C.get("ready"):
        raise HTTPException(503, "Initialisation en cours")
    try:
        chunks = C["retriever"].retrieve(query=q, skip_cache=True, top_k=top_k)
        return {"query": q, "results": [_chunk_to_dict(c) for c in chunks]}
    except Exception as e:
        raise HTTPException(500, f"Recherche indisponible : {e}")


# ═══════════════════════════════════════════════════════════════════════
#  UPLOAD DE RÈGLEMENTS — ingestion à la demande (extraction → traduction →
#  chunking → embeddings → Weaviate). Le document devient interrogeable par
#  le RAG. Traité en arrière-plan (thread), le front sonde le statut.
# ═══════════════════════════════════════════════════════════════════════
REG_JOBS: dict[str, dict] = {}
_reg_counter = {"n": 0}


class RegulationUploadIn(BaseModel):
    filename:     str
    content_b64:  str = Field(..., description="PDF encodé base64")
    jurisdiction: str = "international"   # "tunisia" | "international"
    translate:    bool = True            # traduire FR→EN (corpus en anglais)


def _run_ingest(job_id: str, tmp_path: str, jurisdiction: str,
                doc_name: str, translate: bool):
    job = REG_JOBS[job_id]

    def progress(stage: str, detail: str = ""):
        job["stage"], job["detail"] = stage, detail

    try:
        from ingest_document import ingest_pdf
        result = ingest_pdf(tmp_path, jurisdiction=jurisdiction,
                            doc_name=doc_name, translate=translate,
                            progress=progress)
        job.update(status="DONE", stage="done", result=result)
    except Exception as e:
        job.update(status="ERROR", stage="error", error=str(e)[:300])
    finally:
        try:
            os.remove(tmp_path)
        except Exception:
            pass


@app.post("/api/regulations/upload")
def regulations_upload(body: RegulationUploadIn,
                       authorization: Optional[str] = Header(None)):
    """Reçoit un PDF réglementaire (base64), lance l'ingestion en tâche de
    fond, renvoie un job_id à sonder via GET /api/regulations/upload/{id}."""
    import base64, tempfile
    if not body.filename.lower().endswith(".pdf"):
        raise HTTPException(400, "Seuls les fichiers PDF sont acceptés")
    try:
        data = base64.b64decode(body.content_b64.split(",")[-1])
    except Exception:
        raise HTTPException(400, "Contenu base64 invalide")
    if len(data) > 40 * 1024 * 1024:
        raise HTTPException(413, "PDF trop volumineux (max 40 Mo)")

    fd = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False)
    fd.write(data)
    fd.close()

    _reg_counter["n"] += 1
    job_id = f"REG-{datetime.now():%Y%m%d}-{_reg_counter['n']:04d}"
    jur = body.jurisdiction if body.jurisdiction in ("tunisia", "international") \
        else "international"
    REG_JOBS[job_id] = {
        "job_id": job_id, "status": "RUNNING", "stage": "queued", "detail": "",
        "doc_name": body.filename, "jurisdiction": jur,
        "created_at": datetime.now().isoformat(),
    }
    threading.Thread(target=_run_ingest,
                     args=(job_id, fd.name, jur, body.filename, body.translate),
                     daemon=True).start()
    return {"job_id": job_id, "status": "RUNNING"}


@app.get("/api/regulations/upload/{job_id}")
def regulations_upload_status(job_id: str):
    job = REG_JOBS.get(job_id)
    if not job:
        raise HTTPException(404, "Job d'ingestion inconnu")
    return job


# ═══════════════════════════════════════════════════════════════════════
#  SIMULATEUR — joue le rôle du FLUX BANCAIRE (source "Core Banking" simulée)
# ═══════════════════════════════════════════════════════════════════════
#  • Génère une transaction toutes les N secondes et la fait passer par le
#    pipeline complet (source=SIMULATOR → visible côté espace admin).
#  • SÉQUENTIEL : le pipeline (~50s) > intervalle → on attend la fin d'une
#    analyse avant la suivante (évite les collisions Weaviate/Neo4j).
#    Intervalle effectif = max(interval, durée pipeline).
#  • RÉUTILISE les mêmes clients au fil des ticks → l'historique Redis
#    s'accumule réellement → profils comportementaux vrais, et le pattern
#    SMURFING ÉMERGE en direct (3+ dépôts <5000 TND/72h du même client).

SIM = {"running": False, "interval": 30, "count": 0,
       "last_tx_id": None, "started_at": None}

# Pool de clients : les 4 seedés (profils connus) + clients simulés (DEFAULT)
_SIM_CLIENTS = ["802559360", "100428660", "100428780", "801176C70",
                "SIM-RETAIL-01", "SIM-RETAIL-02", "SIM-BIZ-01"]
_SIM_RECENT: list = []   # biais de répétition (même client → historique)


def _gen_sim_tx() -> dict:
    """Transaction réaliste — mix de scénarios AML (normal / seuil / smurfing / route à risque)."""
    # 60% : re-piocher un client récent (construit son historique)
    if _SIM_RECENT and random.random() < 0.6:
        client = random.choice(_SIM_RECENT[-5:])
    else:
        client = random.choice(_SIM_CLIENTS)
    _SIM_RECENT.append(client)

    roll = random.random()
    now_iso = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    if roll < 0.55:      # virement local normal
        tx = {"tx_type": "VIREMENT_LOCAL",
              "amount": round(random.uniform(400, 8000), 2), "currency": "TND",
              "sender_country": "Tunisia", "destination_country": "Tunisia",
              "beneficiary_name": f"Account-BEN{random.randint(100, 999)}",
              "motive": "Paiement facture", "has_justification": True}
    elif roll < 0.70:    # SWIFT export
        tx = {"tx_type": "SWIFT",
              "amount": round(random.uniform(3000, 30000), 2), "currency": "USD",
              "sender_country": "Tunisia",
              "destination_country": random.choice(["France", "Germany", "Turkey", "UAE"]),
              "beneficiary_name": f"Account-INT{random.randint(100, 999)}",
              "motive": "Export marchandises",
              "has_justification": random.random() < 0.8}
    elif roll < 0.90:    # dépôt espèces SOUS le seuil (répété → smurfing émergent)
        client = "SIM-RETAIL-01" if random.random() < 0.5 else client
        _SIM_RECENT.append(client)
        tx = {"tx_type": "DEPOT_ESPECES",
              "amount": round(random.uniform(3800, 4950), 2), "currency": "TND",
              "sender_country": "Tunisia", "destination_country": "Tunisia",
              "beneficiary_name": "Compte propre", "motive": "Dépôt espèces",
              "has_justification": random.random() < 0.4}
    else:                # route géographique à haut risque
        tx = {"tx_type": "SWIFT",
              "amount": round(random.uniform(12000, 60000), 2), "currency": "USD",
              "sender_country": random.choice(["Tunisia", "Congo"]),
              "destination_country": random.choice(["Iran", "Syria", "North Korea"]),
              "beneficiary_name": f"Account-HR{random.randint(100, 999)}",
              "motive": "", "has_justification": False}

    tx["sender_name"] = f"Account-{client}"
    tx["client_id"] = client
    tx["timestamp"] = now_iso
    return tx


def _sim_loop():
    while SIM["running"]:
        t0 = time.time()
        if not C.get("ready"):
            time.sleep(2)
            continue
        tx = _gen_sim_tx()
        tx_id = _next_tx_id()
        JOBS[tx_id] = {
            "tx_id": tx_id, "status": "RUNNING", "current_step": None,
            "transaction": tx, "steps": {},
            "source": "SIMULATOR", "owner": None,
            "created_at": datetime.now().isoformat(),
        }
        SIM["count"] += 1
        SIM["last_tx_id"] = tx_id
        execute_pipeline(tx, tx_id)   # synchrone → un pipeline à la fois
        # attendre le reste de l'intervalle (réactif au stop)
        deadline = t0 + SIM["interval"]
        while SIM["running"] and time.time() < deadline:
            time.sleep(0.5)


class SimulatorIn(BaseModel):
    interval_seconds: int = Field(30, ge=10, le=600)


@app.post("/api/simulator/start")
def simulator_start(body: SimulatorIn):
    """Démarre le flux bancaire simulé (1 transaction / intervalle, séquentiel)."""
    if not C.get("ready"):
        raise HTTPException(503, "Composants en cours d'initialisation")
    if SIM["running"]:
        return {"running": True, "message": "Simulateur déjà actif",
                "count": SIM["count"]}
    SIM.update(running=True, interval=body.interval_seconds,
               started_at=datetime.now().isoformat())
    threading.Thread(target=_sim_loop, daemon=True).start()
    return {"running": True, "interval_seconds": SIM["interval"],
            "note": "Séquentiel : intervalle effectif = max(intervalle, durée pipeline ~50s)"}


@app.post("/api/simulator/stop")
def simulator_stop():
    SIM["running"] = False
    return {"running": False, "generated": SIM["count"]}


@app.get("/api/simulator/status")
def simulator_status():
    return {"running": SIM["running"], "interval_seconds": SIM["interval"],
            "generated": SIM["count"], "last_tx_id": SIM["last_tx_id"],
            "started_at": SIM["started_at"],
            "mode": SIM.get("mode"),
            "speedrun": SIM.get("speedrun", [])}


# ── SCÉNARIO RÉALISTE (« journée de banque ») ────────────────────────────
#  Rejoue une séquence de transactions COMME SI de vrais clients les
#  soumettaient : chaque opération est rattachée à un VRAI compte utilisateur
#  (owner = email → apparaît dans « Mes dossiers » du client, et la décision
#  du superviseur déclenche le VRAI email vert/rouge), certaines joignent un
#  JUSTIFICATIF réel (PDF/DOCX → extraction docling en étape 0), et le même
#  client fait plusieurs dépôts → l'historique Redis s'accumule → smurfing
#  et LSTM. Toute la palette PASS/EDD/STR/BLOCK en ~5 minutes, zéro saisie.

_SCN_USERS = ["agent@demo.local",
              "agent2@demo.local",
              "agent3@demo.local"]

REALISTIC_SCENARIOS = [
    {"owner": _SCN_USERS[0], "expected": "PASS",
     "doc": "recu_paiement_2500_tnd.txt",
     "tx": {"tx_type": "VIREMENT_LOCAL", "amount": 2500, "currency": "TND",
            "sender_country": "Tunisia", "destination_country": "Tunisia",
            "beneficiary_name": "Ste Services Plus",
            "motive": "Paiement facture électricité", "has_justification": True}},
    {"owner": _SCN_USERS[1], "expected": "EDD",
     "doc": "facture_export_12000_usd.pdf",
     "tx": {"tx_type": "SWIFT", "amount": 12000, "currency": "USD",
            "sender_country": "Tunisia", "destination_country": "France",
            "beneficiary_name": "Societe Lumiere SARL",
            "motive": "Paiement facture export", "has_justification": True}},
    {"owner": _SCN_USERS[2], "expected": "STR",
     "doc": "contrat_prestation_15000_tnd.docx",
     "tx": {"tx_type": "VIREMENT_LOCAL", "amount": 15000, "currency": "TND",
            "sender_country": "Tunisia", "destination_country": "Tunisia",
            "beneficiary_name": "Bureau Etudes Manar SARL",
            "motive": "Acompte contrat prestation", "has_justification": True}},
    {"owner": _SCN_USERS[1], "expected": "STR", "doc": None,
     "tx": {"tx_type": "SWIFT", "amount": 18000, "currency": "USD",
            "sender_country": "Tunisia",
            "destination_country": "United Arab Emirates",
            "beneficiary_name": "Gulf Trading FZE",
            "motive": "", "has_justification": False}},
    {"owner": _SCN_USERS[0], "expected": "PASS", "doc": None,
     "tx": {"tx_type": "DEPOT_ESPECES", "amount": 4800, "currency": "TND",
            "sender_country": "Tunisia", "destination_country": "Tunisia",
            "beneficiary_name": "Compte propre",
            "motive": "Dépôt recettes commerce", "has_justification": True}},
    {"owner": _SCN_USERS[0], "expected": "PASS", "doc": None,
     "tx": {"tx_type": "DEPOT_ESPECES", "amount": 4500, "currency": "TND",
            "sender_country": "Tunisia", "destination_country": "Tunisia",
            "beneficiary_name": "Compte propre",
            "motive": "Dépôt recettes commerce", "has_justification": True}},
    {"owner": _SCN_USERS[0], "expected": "STR", "doc": None,   # 3e dépôt → smurfing
     "tx": {"tx_type": "DEPOT_ESPECES", "amount": 4900, "currency": "TND",
            "sender_country": "Tunisia", "destination_country": "Tunisia",
            "beneficiary_name": "Compte propre",
            "motive": "Dépôt recettes commerce", "has_justification": True}},
    {"owner": _SCN_USERS[2], "expected": "BLOCK", "doc": None,
     "tx": {"tx_type": "SWIFT", "amount": 20000, "currency": "USD",
            "sender_country": "Tunisia", "destination_country": "Iran",
            "beneficiary_name": "Teheran Import Co",
            "motive": "Import marchandises", "has_justification": False}},
]


def _speedrun_loop():
    import base64
    for i, scn in enumerate(REALISTIC_SCENARIOS, 1):
        if not SIM["running"]:
            break
        owner = scn["owner"]
        tx = dict(scn["tx"])
        tx["sender_name"] = f"Account-{owner}"   # client_id stable = email
        tx["client_id"] = owner
        tx["timestamp"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        files = []
        if scn.get("doc"):
            try:
                data = (BASE / "test_documents" / scn["doc"]).read_bytes()
                files = [{"filename": scn["doc"],
                          "content_b64": base64.b64encode(data).decode()}]
            except Exception as e:
                print(f"  [API] scénario : document {scn['doc']} illisible ({e})")
        tx_id = _next_tx_id()
        JOBS[tx_id] = {
            "tx_id": tx_id, "status": "RUNNING", "current_step": None,
            "transaction": tx, "steps": {}, "_files": files,
            # source USER + owner réel → la transaction apparaît dans
            # « Mes dossiers » du client, et la validation superviseur
            # déclenche le VRAI email de décision.
            "source": "USER", "owner": owner,
            "created_at": datetime.now().isoformat(),
        }
        SIM["count"] += 1
        SIM["last_tx_id"] = tx_id
        entry = {"tx_id": tx_id,
                 "label": f"{tx['tx_type']} {tx['amount']:,.0f} {tx['currency']} "
                          f"· {owner.split('@')[0]}"
                          + (f" · 📎 {scn['doc']}" if scn.get("doc") else ""),
                 "expected": scn["expected"],
                 "got": None, "status": "RUNNING"}
        SIM.setdefault("speedrun", []).append(entry)
        execute_pipeline(tx, tx_id)          # synchrone, un à la fois
        job = JOBS.get(tx_id, {})
        entry["status"] = job.get("status", "?")
        entry["got"] = (job.get("summary") or {}).get("ca_decision")
    SIM.update(running=False, mode=None)


@app.post("/api/simulator/speedrun")
def speedrun_start():
    """Scénario réaliste : 8 opérations de vrais clients (documents + emails)."""
    if not C.get("ready"):
        raise HTTPException(503, "Composants en cours d'initialisation")
    if SIM["running"]:
        raise HTTPException(409, "Un flux simulé tourne déjà — arrête-le d'abord")
    SIM.update(running=True, mode="speedrun", speedrun=[],
               started_at=datetime.now().isoformat())
    threading.Thread(target=_speedrun_loop, daemon=True).start()
    return {"running": True, "mode": "speedrun",
            "note": "8 opérations clients enchaînées (documents + smurfing + emails)"}


# ═══════════════════════════════════════════════════════════════════════
#  DATASET IBM HI-SMALL — stats pré-calculées offline (page admin)
# ═══════════════════════════════════════════════════════════════════════
#  Le fichier output/dataset_stats.json est produit UNE FOIS par
#  compute_dataset_stats.py (5M+ lignes agrégées) — ici on le sert tel quel.

DATASET_STATS_PATH = BASE / "output" / "dataset_stats.json"


@app.get("/api/dataset-stats")
def dataset_stats():
    """Agrégats du dataset IBM AML HI-Small (pré-calculés hors-ligne)."""
    if not DATASET_STATS_PATH.exists():
        raise HTTPException(
            404,
            "Stats non calculées — lancer une fois : python compute_dataset_stats.py")
    return json.loads(DATASET_STATS_PATH.read_text(encoding="utf-8"))


# ═══════════════════════════════════════════════════════════════════════
#  COUCHE CONNAISSANCE — corpus, règles, graphe (pages espace Conformité)
# ═══════════════════════════════════════════════════════════════════════
#  Schéma réel (neo4j_graph/schema.py) : nœuds Document/LegalUnit/Chunk/Rule
#  + référentiels ; les 49 règles ground-truth ont origin='human_verified'
#  (PAS de label :GroundTruth). Relations : FROM_DOCUMENT, UNDER_JURISDICTION,
#  HAS_ROLE, APPLIES_TO, NEXT_IN_DOC, REALIZES, SUPPORTED_BY, DERIVED_FROM.

_GRAPH_DRIVER = None


def _graph():
    """Driver Neo4j paresseux, indépendant de l'init du pipeline."""
    global _GRAPH_DRIVER
    if _GRAPH_DRIVER is None:
        import os
        from neo4j import GraphDatabase
        _GRAPH_DRIVER = GraphDatabase.driver(
            os.getenv("NEO4J_URI", "bolt://localhost:7687"),
            auth=(os.getenv("NEO4J_USER", "neo4j"),
                  os.getenv("NEO4J_PASSWORD", "amlkyc2025")))
        _GRAPH_DRIVER.verify_connectivity()
    return _GRAPH_DRIVER


def _jsafe(v):
    """Convertit les types Neo4j non-JSON (dates…) en chaînes."""
    if isinstance(v, (str, int, float, bool)) or v is None:
        return v
    if isinstance(v, list):
        return [_jsafe(x) for x in v]
    return str(v)


def _normalize_rule(p: dict) -> dict:
    """Uniformise les propriétés d'un nœud Rule vers un contrat FRONT lisible.

    Le graphe stocke : component_subject / component_constraint (le contenu de
    la règle), chunk_role (la catégorie : OBLIGATION/KYC/THRESHOLD/…),
    source_article, source_text_snippet (citation verbatim). La page « Règles
    extraites » attendait text/category/article_ref/source_snippet → on mappe.
    """
    p = {k: _jsafe(v) for k, v in p.items()}
    subject = (p.get("component_subject") or "").strip()
    constraint = (p.get("component_constraint") or "").strip()
    text = (p.get("rule_text") or p.get("text") or "").strip()
    if not text:
        if subject and constraint:
            text = f"{subject} — {constraint}"
        else:
            text = subject or constraint or "(règle sans texte)"
    return {
        **p,
        "id": p.get("id"),
        "text": text,
        "subject": subject,
        "constraint": constraint,
        "category": p.get("chunk_role") or p.get("category") or "",
        "article_ref": p.get("source_article") or p.get("article_ref") or "",
        "source_doc": p.get("source_doc") or "",
        "source_snippet": p.get("source_text_snippet") or p.get("source_snippet") or "",
        "jurisdiction": p.get("jurisdiction") or "",
        "confidence": p.get("confidence"),
        "origin": p.get("origin") or p.get("status") or "inconnu",
    }


@app.get("/api/corpus/documents")
def corpus_documents():
    """Bibliothèque du corpus : documents sources + comptes réels (Neo4j)."""
    try:
        with _graph().session() as s:
            rows = s.run("""
                MATCH (d:Document)
                OPTIONAL MATCH (lu:LegalUnit)-[:FROM_DOCUMENT]->(d)
                WITH d, count(DISTINCT lu) AS legal_units
                OPTIONAL MATCH (c:Chunk {source_doc: d.name})
                RETURN d.name AS name, d.jurisdiction AS jurisdiction,
                       d.priority AS priority, legal_units,
                       count(c) AS chunks
                ORDER BY chunks DESC
            """).data()
        return {"documents": rows,
                "total_documents": len(rows),
                "total_chunks": sum(r["chunks"] for r in rows),
                "jurisdictions": sorted({r["jurisdiction"] for r in rows
                                         if r["jurisdiction"]})}
    except Exception as e:
        raise HTTPException(503, f"Neo4j indisponible : {e}")


@app.get("/api/corpus/chunks")
def corpus_chunks(doc: str, limit: int = 40):
    """Structure d'un document : ses chunks (métadonnées — le texte intégral
    vit dans Weaviate, le graphe garde la structure articulaire)."""
    try:
        with _graph().session() as s:
            rows = s.run("""
                MATCH (c:Chunk {source_doc: $doc})
                OPTIONAL MATCH (c)-[:REALIZES]->(lu:LegalUnit)
                RETURN c.chunk_id AS chunk_id, c.article_ref AS article_ref,
                       c.chunk_role AS chunk_role,
                       c.has_threshold AS has_threshold,
                       lu.primary_role AS unit_role
                ORDER BY c.chunk_id
                LIMIT $limit
            """, doc=doc, limit=limit).data()
        return {"doc": doc, "chunks": rows}
    except Exception as e:
        raise HTTPException(503, f"Neo4j indisponible : {e}")


# ── Fichiers PDF sources (PFE_AML_KYC/01_regulations) ────────────────────
#  Chemin surchargeable par la variable d'environnement REG_DIR : en conteneur,
#  le corpus est MONTÉ en volume (68 Mo hors contexte de build) plutôt que copié
#  dans l'image.
REG_DIR = Path(os.environ.get("REG_DIR")
               or (BASE.parent / "PFE_AML_KYC" / "01_regulations")).resolve()


@app.get("/api/corpus/files")
def corpus_files():
    """Liste les PDF réglementaires réels (international/ + tunisia/)."""
    if not REG_DIR.exists():
        raise HTTPException(404, f"Dossier introuvable : {REG_DIR}")
    files = []
    for p in sorted(REG_DIR.rglob("*.pdf")):
        try:
            rel = p.relative_to(REG_DIR).as_posix()
        except ValueError:
            continue
        files.append({
            "name": p.name,
            "folder": p.parent.name,
            "rel_path": rel,
            "size_kb": round(p.stat().st_size / 1024, 1),
        })
    return {"files": files, "total": len(files),
            "folders": sorted({f["folder"] for f in files})}


@app.get("/api/corpus/file")
def corpus_file(path: str):
    """Sert un PDF pour la prévisualisation. `path` = rel_path de /corpus/files.

    Sécurité : on résout le chemin et on vérifie qu'il reste SOUS REG_DIR
    (empêche la traversée de répertoire ../../)."""
    target = (REG_DIR / path).resolve()
    if REG_DIR not in target.parents or target.suffix.lower() != ".pdf":
        raise HTTPException(400, "Chemin invalide")
    if not target.exists():
        raise HTTPException(404, "Fichier introuvable")
    return FileResponse(str(target), media_type="application/pdf",
                        filename=target.name,
                        headers={"Content-Disposition": f'inline; filename="{target.name}"'})


@app.get("/api/rules")
def list_rules(origin: Optional[str] = None, limit: int = 300):
    """Règles du graphe : 49 ground-truth (origin=human_verified) + règles
    apprises par le RAG (agent_extracted/agent_validated ou status LEARNED…)."""
    try:
        with _graph().session() as s:
            rows = s.run("""
                MATCH (r:Rule)
                WHERE $origin IS NULL OR r.origin = $origin OR r.status = $origin
                RETURN properties(r) AS p
                LIMIT $limit
            """, origin=origin, limit=limit).data()
        rules = [_normalize_rule(row["p"]) for row in rows]
        # Ground-truth d'abord, puis par confiance décroissante
        rules.sort(key=lambda r: (r.get("origin") != "human_verified",
                                  -(r.get("confidence") or 0)))
        gt = sum(1 for r in rules if r.get("origin") == "human_verified")
        return {"rules": rules, "total": len(rules),
                "ground_truth": gt, "learned": len(rules) - gt}
    except Exception as e:
        raise HTTPException(503, f"Neo4j indisponible : {e}")


@app.get("/api/rules/{rule_id}")
def rule_detail(rule_id: str):
    """Détail d'une règle + ses preuves (chunks SUPPORTED_BY, LegalUnits
    DERIVED_FROM) — la chaîne d'ancrage anti-hallucination."""
    try:
        with _graph().session() as s:
            row = s.run("MATCH (r:Rule {id: $id}) RETURN properties(r) AS p",
                        id=rule_id).single()
            if not row:
                raise HTTPException(404, "Règle inconnue")
            evidence = s.run("""
                MATCH (r:Rule {id: $id})-[:SUPPORTED_BY]->(c:Chunk)
                OPTIONAL MATCH (c)-[:REALIZES]->(lu:LegalUnit)
                RETURN c.source_doc AS source_doc, c.article_ref AS article_ref,
                       c.chunk_role AS chunk_role, lu.unit_id AS unit_id
                LIMIT 5
            """, id=rule_id).data()
            derived = s.run("""
                MATCH (r:Rule {id: $id})-[:DERIVED_FROM]->(lu:LegalUnit)
                RETURN lu.unit_id AS unit_id, lu.source_doc AS source_doc,
                       lu.article_ref AS article_ref
                LIMIT 5
            """, id=rule_id).data()
        return {"rule": _normalize_rule(row["p"]),
                "evidence_chunks": evidence, "derived_from": derived}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(503, f"Neo4j indisponible : {e}")


@app.get("/api/graph/stats")
def graph_stats():
    """Statistiques du graphe de connaissances (nœuds/label, relations/type)."""
    try:
        with _graph().session() as s:
            nodes = {r["label"]: r["count"] for r in s.run(
                "MATCH (n) RETURN labels(n)[0] AS label, count(*) AS count "
                "ORDER BY count DESC").data()}
            rels = {r["type"]: r["count"] for r in s.run(
                "MATCH ()-[x]->() RETURN type(x) AS type, count(*) AS count "
                "ORDER BY count DESC").data()}
        return {"nodes": nodes, "relationships": rels,
                "total_nodes": sum(nodes.values()),
                "total_relationships": sum(rels.values())}
    except Exception as e:
        raise HTTPException(503, f"Neo4j indisponible : {e}")


@app.get("/api/graph/sample")
def graph_sample(limit: int = 8):
    """Sous-graphe RÉEL pour visualisation directe dans la plateforme (sans
    dépendre de Neo4j Browser, qui refuse l'iframe). Renvoie la chaîne
    d'ancrage anti-hallucination : Rule → Chunk → LegalUnit → Document."""
    try:
        with _graph().session() as s:
            rows = s.run("""
                MATCH (r:Rule {origin:'human_verified'})-[:SUPPORTED_BY]->(c:Chunk)
                      -[:REALIZES]->(lu:LegalUnit)-[:FROM_DOCUMENT]->(d:Document)
                WITH r, c, lu, d LIMIT $lim
                RETURN r.id AS r_id,
                       coalesce(r.chunk_role, 'RULE') AS r_cap,
                       c.chunk_id AS c_id, c.article_ref AS c_cap,
                       lu.unit_id AS lu_id,
                       coalesce(lu.article_ref, lu.unit_id) AS lu_cap,
                       d.name AS d_id, d.jurisdiction AS d_juris
            """, lim=limit).data()
        nodes: dict = {}
        edges: list = []
        seen = set()

        def add(nid, ntype, cap):
            if nid and nid not in nodes:
                nodes[nid] = {"id": nid, "type": ntype,
                              "caption": str(cap or nid)[:48]}

        def link(a, b, t):
            k = (a, b, t)
            if a and b and k not in seen:
                seen.add(k)
                edges.append({"source": a, "target": b, "type": t})

        for row in rows:
            add(row["r_id"], "Rule", row["r_id"])
            add(row["c_id"], "Chunk", row["c_cap"])
            add(row["lu_id"], "LegalUnit", row["lu_cap"])
            add(row["d_id"], "Document", str(row["d_id"]).replace(".pdf", ""))
            link(row["r_id"], row["c_id"], "SUPPORTED_BY")
            link(row["c_id"], row["lu_id"], "REALIZES")
            link(row["lu_id"], row["d_id"], "FROM_DOCUMENT")
        return {"nodes": list(nodes.values()), "edges": edges,
                "counts": {"nodes": len(nodes), "edges": len(edges)}}
    except Exception as e:
        raise HTTPException(503, f"Neo4j indisponible : {e}")
